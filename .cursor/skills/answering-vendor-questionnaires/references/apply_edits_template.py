"""Reusable harness for applying answer edits to a vendor security
questionnaire .docx. Copy this file into /tmp and adjust the EDITS list
before running.

Pattern:

  - Backs up the source docx to a sibling *.pre_edit.docx the first time it
    runs.
  - Pre-flights every edit: confirms the `old_string` substring is present
    exactly once in the target paragraph (idempotent re-runs are skipped
    if `new_string` is already present and `old_string` is gone).
  - Walks each EDIT, applies the substring replacement to the response
    cell's answer paragraph (table row 1, column 0, paragraph index 1),
    preserving the bold "Response:" label paragraph at index 0.
  - Saves, reloads, and verifies that every `old_string` is gone and every
    `new_string` is present. Fails loud on mismatch.

Notes:

  - The response cell layout is two paragraphs:
      P0 = bold "Response:" label, single run
      P1 = answer, typically a single non-bold run
  - Substring replacement on `paragraphs[1].runs[0].text` preserves all
    formatting. Whole-paragraph rewrites should set `runs[0].text =
    new_text` and leave the run's other formatting attributes alone.
  - For full-paragraph rewrites, set EDITS[i]["old"] to the entire
    existing answer text and EDITS[i]["new"] to the entire new answer.
"""

import shutil
from pathlib import Path

from docx import Document

DOCX = Path("/path/to/Vendor_Security_Questionnaire_FILLED.docx")
BACKUP = DOCX.with_suffix(".pre_edit.docx")

EDITS = [
    # Each entry targets one response cell by table index. The QID column is
    # for chat output only -- the actual lookup uses tbl_idx.
    # {
    #     "tbl_idx": 29,
    #     "qid": "3.1",
    #     "old": "exact substring to replace",
    #     "new": "replacement text",
    # },
]


def main() -> None:
    if not BACKUP.exists():
        shutil.copy2(DOCX, BACKUP)
        print(f"[backup] wrote {BACKUP}")
    else:
        print(f"[backup] {BACKUP} already exists -- leaving in place")

    doc = Document(str(DOCX))

    for e in EDITS:
        tbl = doc.tables[e["tbl_idx"]]
        p1 = tbl.rows[1].cells[0].paragraphs[1]
        text = p1.runs[0].text
        if e["new"] in text and e["old"] not in text:
            print(f"[skip] {e['qid']} already patched")
            e["_skip"] = True
            continue
        count = text.count(e["old"])
        if count != 1:
            raise SystemExit(
                f"[fatal] {e['qid']}: expected 1 occurrence of "
                f"{e['old']!r}, found {count}. Aborting before any write."
            )

    for e in EDITS:
        if e.get("_skip"):
            continue
        tbl = doc.tables[e["tbl_idx"]]
        p1 = tbl.rows[1].cells[0].paragraphs[1]
        run = p1.runs[0]
        before = run.text
        after = before.replace(e["old"], e["new"], 1)
        if after == before:
            raise SystemExit(
                f"[fatal] {e['qid']}: replacement was a no-op."
            )
        run.text = after
        print(f"[edit] {e['qid']}: applied ({len(before)} -> {len(after)} chars)")

    doc.save(str(DOCX))
    print(f"[save] wrote {DOCX}")

    doc2 = Document(str(DOCX))
    for e in EDITS:
        tbl = doc2.tables[e["tbl_idx"]]
        p1 = tbl.rows[1].cells[0].paragraphs[1]
        text = p1.runs[0].text
        if e["old"] in text:
            raise SystemExit(
                f"[fatal] {e['qid']}: old substring still present after save."
            )
        if e["new"] not in text:
            raise SystemExit(
                f"[fatal] {e['qid']}: new substring not found after save."
            )
        print(f"[verify] {e['qid']}: OK")


if __name__ == "__main__":
    main()
